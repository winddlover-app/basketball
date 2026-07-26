from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"C:\Users\tonyf\OneDrive\篮球APP\Basketball_Training_Camp_APP_PRD_EN.docx"
FONT = "Calibri"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def para_spacing(paragraph, before=0, after=6, line=1.1):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:cantSplit"))
    if node is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        keep_row(row)
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width * 1440)))


def title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_spacing(p, 0, 8, 1.1)
    r = p.add_run("Basketball Training Camp APP Product Requirements Document")
    set_run_font(r, 21, True, "0B2545")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_spacing(p, 0, 18, 1.1)
    r = p.add_run("Course operations, after-class training, video analysis, and coach feedback for basketball training organizations")
    set_run_font(r, 10.5, False, "4A5568")

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Document Type", "Product Requirements Document (PRD)"),
        ("Audience", "Product, design, operations, business owners, QA, and planning stakeholders"),
        ("Version", "V1.0 - English Interface Edition"),
        ("Date", "2026-07-26"),
    ]
    for i, (label, value) in enumerate(rows):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
        shade(table.rows[i].cells[0], "F2F4F7")
    set_table_width(table, [1.55, 4.95])
    format_table(table)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.text = text
    sizes = {1: 16, 2: 13, 3: 12}
    colors = {1: "2E74B5", 2: "2E74B5", 3: "1F4D78"}
    para_spacing(p, {1: 16, 2: 12, 3: 8}.get(level, 6), {1: 8, 2: 6, 3: 4}.get(level, 4), 1.1)
    for run in p.runs:
        set_run_font(run, sizes.get(level, 11), True, colors.get(level, "111827"))


def body(doc, text):
    p = doc.add_paragraph()
    para_spacing(p, 0, 5, 1.08)
    r = p.add_run(text)
    set_run_font(r, 11)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    para_spacing(p, 0, 3, 1.12)
    r = p.add_run(text)
    set_run_font(r, 11)


def num(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    para_spacing(p, 0, 3, 1.12)
    r = p.add_run(text)
    set_run_font(r, 11)


def format_table(table):
    for ri, row in enumerate(table.rows):
        keep_row(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                para_spacing(p, 0, 0, 1.0)
                for run in p.runs:
                    set_run_font(run, 9.0, ri == 0, "111827")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        shade(table.rows[0].cells[i], "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_width(table, widths)
    format_table(table)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"]:
    style = doc.styles[style_name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(11)

title(doc)

heading(doc, "1. Product Overview")
body(doc, "The Basketball Training Camp APP serves basketball training organizations, coaches, students, and parents. It brings course enrollment, class operations, after-class training, training video analysis, coach feedback, and growth records into one continuous service flow.")
body(doc, "The product focuses on solving common offline training pain points: fragmented service information, limited visibility into after-class practice, hard-to-preserve video feedback, and inconsistent coach evaluation records.")

heading(doc, "2. Product Goals")
for item in [
    "Digitize course services: students and parents can browse courses, enroll, pay, view schedules, and track course status.",
    "Make after-class training trackable: coaches can assign tasks, students can submit videos, and both sides can follow task progress.",
    "Support training video analysis: the system can generate preliminary metrics, key moments, and improvement suggestions from submitted videos.",
    "Preserve coach feedback: coaches review and publish official feedback reports that students and parents can revisit over time.",
    "Standardize institutional service delivery: unified workflows, states, and reports help improve service quality and renewal conversion.",
]:
    bullet(doc, item)

heading(doc, "3. User Roles and Core Value")
add_table(doc, ["Role", "Core Need", "Product Value"], [
    ["Student", "Know what to practice, how to practice, and what needs improvement.", "View tasks, upload training videos, receive feedback, and track growth records."],
    ["Parent", "Understand whether the child is attending class and improving through training.", "View courses, orders, feedback reports, and stage-based growth records."],
    ["Coach", "Reduce after-class follow-up effort and improve feedback efficiency.", "Take attendance, assign tasks, review videos, validate AI reports, and publish evaluations."],
    ["Institution Operator", "Standardize service delivery and improve renewal conversion.", "Configure courses, students, coaches, payments, videos, and report status."],
    ["Super Admin", "Manage user access and global business records.", "Approve parent/coach registration, manage user lists, and trace courses and feedback records."],
], [1.15, 2.45, 2.9])

heading(doc, "4. Service Flow")
for item in [
    "Student/Parent browses courses and filters by campus, age, level, date, and coach.",
    "Student/Parent submits enrollment and completes payment; the system creates a personal schedule and session records.",
    "Coach receives today's courses and student roster, then checks class information and safety notes before class.",
    "Coach takes attendance and records Present, Late, Leave Requested, or Absent.",
    "Coach assigns after-class training tasks to an individual student or the full class, including action, repetitions, demo video, deadline, and completion criteria.",
    "Student records or uploads a training video based on the assigned task and views upload/analysis status.",
    "The system generates preliminary training metrics, key clips, and correction suggestions.",
    "Coach reviews the result, confirms, edits, or hides suggestions, then publishes the official report.",
    "Student/Parent views the feedback report; the system preserves training history and growth trends.",
]:
    num(doc, item)

heading(doc, "5. Scope Plan")
heading(doc, "5.1 Phase 1 Scope", 2)
for item in [
    "Account and profile foundation: registration/login, role recognition and switching, profile management, privacy authorization, guardian binding, and super admin user management.",
    "Student/Parent course flow: course browsing, course details, enrollment, payment, order/class-hour records, schedule, training feedback, and leave/reschedule/makeup requests.",
    "Coach class and feedback flow: today's courses, weekly schedule, student roster, safety notes, attendance, AI report review, training plan, coach demo videos, training tasks, and history records.",
    "AI-assisted video action analysis: record/import video, choose training item, shooting guidance, upload status, basic shooting and dribbling recognition, body keypoints/angles, text suggestions, and report/clip storage.",
    "Notifications: send pre-class reminders, reschedule notices, and cancellation notices by email.",
]:
    bullet(doc, item)

heading(doc, "5.2 Phase 2 Scope", 2)
for item in [
    "Student/Parent schedule enhancement: personal schedule and session details with list/calendar views.",
    "Coach task enhancement: task assignment, submission filtering, video review, task acceptance, rejection, and evaluation.",
    "AI enhancement: real-time low-latency analysis as a future capability, evaluated separately before launch.",
    "Service support enhancement: contact support or submit feedback with issue details, contact method, screenshots, and processing status.",
]:
    bullet(doc, item)

heading(doc, "6. Phase 1 Functional Requirements")
heading(doc, "6.1 Common Account and Basic Features", 2)
add_table(doc, ["ID", "Requirement", "Acceptance Criteria"], [
    ["APP-001", "Email/password registration, login, and logout.", "Verification code expiry, rate limiting, and error messages are clear; account state is correct after logout."],
    ["APP-002", "Recognize and switch between Student, Parent, and Coach roles.", "Each role only sees allowed features and data; page content refreshes after role switching."],
    ["APP-003", "Maintain personal profile.", "Supports avatar, name, gender, birthday, and contact details; required fields show clear prompts; changes sync across related pages."],
    ["APP-004", "Bind guardians for minors.", "When a minor enrolls, the system can match a parent email; matching email and birth date allow binding; parents can switch between multiple children."],
    ["APP-005", "Privacy Policy, Terms of Use, and Guardian Authorization confirmation.", "Records consent version and time; video submission and AI analysis are restricted before required authorization."],
    ["APP-006", "Super Admin manages parent and coach registration/cancellation.", "Supports list-based user management and tracing of linked students, courses, account duration, course history, and coach feedback."],
], [0.85, 2.45, 3.2])

heading(doc, "6.2 Student/Parent App", 2)
add_table(doc, ["ID", "Requirement", "Acceptance Criteria"], [
    ["STU-001", "Browse and filter course list.", "Filters include campus, age, level, date, and coach; list shows price, remaining seats, time, age range, and status."],
    ["STU-002", "View course details and submit enrollment.", "Validates age, capacity, and eligibility; duplicate enrollment gives a clear prompt."],
    ["STU-003", "Pay course fee online and view payment result.", "After successful payment, order, seats, class hours, and schedule are updated; failed payments can be retried."],
    ["STU-007", "View coach rating, text evaluation, AI report, and training history.", "Reports are linked to course or task; filtering by time and training item is supported."],
    ["STU-008", "Request leave, rescheduling, or makeup classes.", "Rules and processing status are shown; final result syncs to the personal schedule."],
    ["STU-009", "Share stage training report with a designated guardian.", "Sharing can be cancelled; report includes coach evaluation, AI metrics, and trends."],
], [0.85, 2.45, 3.2])

heading(doc, "6.3 Coach App", 2)
add_table(doc, ["ID", "Requirement", "Acceptance Criteria"], [
    ["COA-001", "View today's courses, weekly schedule, and course details.", "Only self-owned or authorized courses are visible; changes sync in time."],
    ["COA-002", "View class student roster, basic profile, and safety notes.", "Sensitive information is restricted; only authorized students are visible."],
    ["COA-003", "Take class attendance.", "Supports Present, Late, Leave Requested, and Absent; batch operation is supported; saved results are traceable."],
    ["COA-007", "Review AI reports and confirm, edit, or hide suggestions.", "Coach confirmation is required before publishing; modification history is retained."],
    ["COA-008", "Create and adjust personal training plans.", "Templates can be referenced; changes sync to the student side."],
    ["COA-009", "Record coach demonstration videos quickly.", "Videos can be linked to tasks or plans; retake and delete are supported."],
    ["COA-010", "View training tasks, video upload status, analysis status, feedback, and history.", "Coach can view the full training service record by student, course, or task."],
], [0.85, 2.45, 3.2])

heading(doc, "6.4 AI-Assisted Video Action Analysis", 2)
add_table(doc, ["ID", "Requirement", "Acceptance Criteria"], [
    ["AI-001", "Record or import training videos.", "Supports start, pause, finish, and retake; camera and album permission prompts are clear."],
    ["AI-002", "Select training item before shooting.", "Supports spot shooting, free throws, and stationary/basic dribbling; each item has matching shooting guidance and analysis metrics."],
    ["AI-003", "Provide shooting guidance.", "Prompts include camera position, distance, lighting, full-body framing, and recording area; adjustment advice is shown when conditions are not met."],
    ["AI-004", "Upload videos and retry failed uploads.", "Weak network handling continues where possible; failure reasons are shown; the same submission should not create duplicate records."],
    ["AI-005", "Recognize shooting actions.", "Shows attempts, made/missed results, shooting percentage, and key timestamps; low-confidence results allow manual correction."],
    ["AI-006", "Recognize dribbling actions.", "Shows count, left/right hand, and rhythm; results can be viewed by clip and manually corrected."],
    ["AI-007", "Show body keypoints and basic angle metrics.", "Metrics link to key frames; recognition failure reasons are shown when applicable."],
    ["AI-008", "Generate text correction suggestions and recommended drills.", "Suggestions are specific, actionable, linked to problem clips, and published only after coach review."],
    ["AI-009", "Save AI analysis reports and key video clips.", "Reports can be revisited; deleting a video prompts users about the impact on reports."],
    ["AI-011", "Recognize ball-control duration and gaps between dribbling actions.", "Shows possession duration within valid time and the gap before the next action."],
], [0.85, 2.45, 3.2])

heading(doc, "6.5 Notifications and Service", 2)
add_table(doc, ["ID", "Requirement", "Acceptance Criteria"], [
    ["MSG-002", "Receive pre-class reminders, reschedule notices, and cancellation notices.", "All notifications are sent by email and include course, time, location, and change details."],
], [0.85, 2.45, 3.2])

heading(doc, "7. Phase 2 Functional Requirements")
add_table(doc, ["Module", "ID", "Requirement", "Acceptance Criteria"], [
    ["Student/Parent App", "STU-005", "View personal schedule and session details.", "Supports list/calendar view; shows time, location, coach, and status."],
    ["Coach App", "COA-004", "Assign training tasks to an individual or class.", "Can set action, repetitions, demo video, deadline, and completion criteria."],
    ["Coach App", "COA-005", "View task submissions, training videos, and AI analysis status.", "Can filter by Not Submitted, Pending Analysis, Pending Review, and Completed."],
    ["Coach App", "COA-006", "Accept, reject, and evaluate tasks.", "Rejection requires a reason; student receives notification."],
    ["AI Assistance", "AI-010", "Real-time low-latency analysis.", "Future capability; launch depends on separate assessment of scenarios, accuracy, and cost-benefit."],
    ["Messages & Service", "MSG-005", "Contact support or submit issue feedback.", "User can enter issue, contact details, and screenshots; processing status is displayed."],
], [1.2, 0.75, 2.25, 2.3])

heading(doc, "8. Key Screen List")
add_table(doc, ["Area", "Screen / Module", "Description"], [
    ["Common", "Login/Register, Role Selection, Privacy Authorization, Profile", "Shared entry points for all roles."],
    ["Student/Parent", "Home, Course List, Course Details, Enrollment Confirmation, Payment Result, Orders/Class Hours, Personal Schedule", "Completes the course purchase flow."],
    ["Student/Parent", "Training Tasks, Video Upload, Upload/Analysis Status, Training Feedback, History", "Completes the after-class training flow."],
    ["Coach", "Coach Home, Today's Courses, Student Roster, Attendance, Task Assignment, Submission List, Video Review", "Highlights pending coach actions."],
    ["Coach", "AI Report Review, Evaluation Publishing, Training Record View", "Human confirmation point before AI results are officially released."],
    ["Messages", "Message Center, Notification Details, Settings", "All messages are sent by email; no in-app push reminder is required."],
    ["Institution/Admin", "User Management, Course Management, Enrollment Orders, Schedule Management, Feedback Report Management", "Supports institution operations and super admin management."],
], [1.05, 2.75, 2.7])

heading(doc, "9. Core Business Rules")
heading(doc, "9.1 Enrollment and Schedule", 2)
for item in [
    "Course information must include campus, time, location, coach, age range, level requirement, price, remaining seats, and status.",
    "Enrollment must validate age, course capacity, eligibility, and duplicate enrollment.",
    "After successful payment, order, seats, class hours, and personal schedule must be updated.",
    "Leave, reschedule, and makeup requests must show rules, processing status, and final results.",
]:
    bullet(doc, item)

heading(doc, "9.2 Guardian and Authorization", 2)
for item in [
    "Minor students must complete guardian binding before parents can view courses, orders, feedback, and growth records.",
    "Guardian binding can be matched by parent email and student birth date.",
    "Privacy Policy, Terms of Use, and Guardian Authorization must record consent version and consent time.",
    "Video submission and AI analysis functions must be restricted before required authorization is completed.",
]:
    bullet(doc, item)

heading(doc, "9.3 Training Tasks and Videos", 2)
for item in [
    "Training tasks must define action, repetitions, demo video, deadline, and completion criteria.",
    "Students must select a training item before submitting a video and receive matching shooting guidance.",
    "Failed video uploads must show a reason and support retry; the same submission should not duplicate records.",
    "Training videos, AI reports, and coach evaluations must be linked to a course or task.",
]:
    bullet(doc, item)

heading(doc, "9.4 AI Report and Coach Review", 2)
for item in [
    "AI output is preliminary data and suggestions only; it is not published to students/parents as official feedback by default.",
    "Coach review is required before publishing, with options to confirm, edit, or hide suggestions.",
    "Low-confidence, failed, or potentially misleading results must show reasons and support manual correction.",
    "Coach modification history must be retained for service quality tracing.",
]:
    bullet(doc, item)

heading(doc, "10. Key State Definitions")
add_table(doc, ["Object", "States", "Description"], [
    ["Course", "Not Started, Open for Enrollment, Full, In Progress, Ended, Cancelled", "Used in course list, course details, schedule, and notifications."],
    ["Order", "Pending Payment, Paid, Payment Failed, Cancelled, Refunded", "Used for parent/student enrollment and fee results."],
    ["Session", "Upcoming, Present, Late, Leave Requested, Absent, Makeup Completed", "Used for attendance, schedule, and class-hour records."],
    ["Training Task", "Not Started, Pending Submission, Submitted, Pending Analysis, Pending Review, Completed, Rejected", "Used for student submission and coach handling."],
    ["Video", "Pending Upload, Uploading, Upload Failed, Uploaded, Analyzing, Analysis Failed, Analysis Completed", "Used for upload and analysis progress display."],
    ["Feedback Report", "Pending Review, Confirmed, Edited, Suggestions Hidden, Published", "Used for coach review and student/parent viewing."],
], [1.15, 2.35, 3.0])

heading(doc, "11. Feedback Report Content Requirements")
for item in [
    "Training basics: student, course or task, training item, submission time, coach, and report status.",
    "Video result summary: key clips, shooting or dribbling metrics, recognizable action sections, and abnormal prompts.",
    "AI preliminary suggestions: problem description, related clip, recommended correction, and suggested drills.",
    "Official coach feedback: coach rating, text evaluation, adjusted training advice, and whether re-practice or resubmission is required.",
    "Growth trend: history by time and training item to help students and parents understand stage-based progress.",
]:
    bullet(doc, item)

heading(doc, "12. Experience and Compliance Requirements")
for item in [
    "Critical flow prompts must be clear, especially enrollment failure, payment failure, video upload failure, recognition failure, and report pending review.",
    "Each role can only view related courses, students, orders, videos, and reports.",
    "Functions involving minors, guardian authorization, video upload, and AI analysis require clear authorization before use.",
    "AI suggestions must avoid absolute wording; official feedback is based on coach-reviewed published content.",
    "Notifications are sent by email and must include course, time, location, and change details.",
]:
    bullet(doc, item)

heading(doc, "13. Success Metrics")
for item in [
    "Course enrollment conversion: conversion from course detail view to enrollment submission and successful payment.",
    "After-class training completion: task submission rate, on-time submission rate, and resubmission rate after rejection.",
    "Coach handling efficiency: pending review task count, report review duration, and task completion cycle time.",
    "Parent/student engagement: feedback report view rate, stage report share rate, and history record access.",
    "Institution service records: student course history completeness, number of stored training reports, and renewal conversion reference data.",
]:
    bullet(doc, item)

heading(doc, "14. Phase 1 Acceptance Standard")
for item in [
    "Student/Parent can complete the full flow from course browsing to enrollment payment and personal schedule viewing.",
    "Coach can view today's courses and student roster, then complete attendance records.",
    "Coach can assign training-related content; student can submit a training video and view processing status.",
    "The system can generate preliminary training data and suggestions; coach can review and publish official feedback.",
    "Student/Parent can view training feedback reports and historical training records.",
    "Super Admin can manage parent/coach access and trace key service records.",
]:
    bullet(doc, item)

heading(doc, "15. Out of Scope for Phase 1")
for item in [
    "Real-time low-latency training analysis.",
    "Complex content community, training content marketplace, or public leaderboard.",
    "Deep management capabilities for tournament organization, team management, or large clubs.",
    "General sports social features unrelated to course services.",
]:
    bullet(doc, item)

footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Basketball Training Camp APP PRD")
set_run_font(r, 9, False, "6B7280")

doc.save(OUT)
print(OUT)
